import argparse
import json
from dataclasses import dataclass
from datetime import date, datetime, timedelta
import calendar
from typing import Dict, List, Sequence, Set

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from ortools.sat.python import cp_model

SHIFT_TYPES = ("D", "E", "N")
REST_TYPES = ("R", "off")
ALL_TYPES = SHIFT_TYPES + REST_TYPES


@dataclass
class Employee:
    name: str
    preferred_shifts: List[str]
    preferred_shift_weights: Dict[str, int]
    min_rest_days: int
    requested_r_dates: List[date]
    previous_last_shift: str = "1off"
    previous_night_count: int | None = None
    max_consecutive_work_days: int = 4
    allow_fifth_day: bool = True


@dataclass
class RuleConfig:
    min_d: int
    max_d: int
    min_e: int
    max_e: int
    min_n: int
    max_n: int
    national_holidays: Set[date] | None = None
    random_seed: int = 42


def parse_date(value: str) -> date:
    return datetime.strptime(value, "%Y-%m-%d").date()


def date_range(start: date, end: date) -> List[date]:
    days: List[date] = []
    cur = start
    while cur <= end:
        days.append(cur)
        cur += timedelta(days=1)
    return days


def build_employees(raw_employees: Sequence[dict]) -> List[Employee]:
    employees: List[Employee] = []
    for item in raw_employees:
        preferred_raw = item.get("preferred_shifts")
        if preferred_raw is None:
            # 向下相容舊格式：單一偏好班別
            preferred_raw = [item.get("preferred_shift")]
        preferred_shifts = [s for s in preferred_raw if s in SHIFT_TYPES]
        if len(preferred_shifts) == 0:
            raise ValueError(f"{item['name']} 的 preferred_shifts 必須包含 D/E/N")
        if len(preferred_shifts) > 2:
            raise ValueError(f"{item['name']} 最多只能設定兩種偏好班別")
        raw_weights = item.get("preferred_shift_weights", {})
        weights: Dict[str, int] = {}
        if len(preferred_shifts) == 1:
            weights[preferred_shifts[0]] = int(raw_weights.get(preferred_shifts[0], 10))
        else:
            # 預設第一個是主偏好、第二個是次偏好
            weights[preferred_shifts[0]] = int(raw_weights.get(preferred_shifts[0], 10))
            weights[preferred_shifts[1]] = int(raw_weights.get(preferred_shifts[1], 6))
        employees.append(
            Employee(
                name=item["name"],
                preferred_shifts=preferred_shifts,
                preferred_shift_weights=weights,
                min_rest_days=item["min_rest_days"],
                requested_r_dates=[parse_date(d) for d in item.get("requested_r_dates", [])],
                previous_last_shift=item.get("previous_last_shift", "1off"),
                previous_night_count=(
                    int(item["previous_night_count"]) if item.get("previous_night_count") not in (None, "") else None
                ),
                max_consecutive_work_days=item.get("max_consecutive_work_days", 4),
                allow_fifth_day=item.get("allow_fifth_day", True),
            )
        )
    return employees


NIGHT_COUNT_MONTH_MIN = 20  # 該曆月月底（含上段）E+N 合計至少此數；可超過。


def format_night_carryover_display(emp: Employee) -> str:
    """範例：12+8 表示上段累計 12 班夜班，月底前至少還需 8 班 E/N（上段+本月≥門檻）；上段已≥門檻則顯示 +0。"""
    if emp.previous_night_count is None:
        return ""
    p = int(emp.previous_night_count)
    need = max(0, NIGHT_COUNT_MONTH_MIN - p)
    return f"{p}+{need}"


def _preflight_staffing(employees: Sequence[Employee], rule: RuleConfig) -> None:
    """無解時使用者難以判斷原因，先做明顯的人力／下限矛盾檢查。"""
    n = len(employees)
    if n == 0:
        raise ValueError("至少需要一位員工。")
    if rule.min_d + rule.min_e + rule.min_n > n:
        s = rule.min_d + rule.min_e + rule.min_n
        raise ValueError(
            f"每日需求下限加總 D+E+N ≥ {rule.min_d}+{rule.min_e}+{rule.min_n}={s} 人，"
            f"大於員工數（{n} 人）。同一天每人只能上一種班，請降低某日下限或增加人力。"
        )
    can_d = sum(1 for emp in employees if "D" in emp.preferred_shifts)
    can_e = sum(1 for emp in employees if "E" in emp.preferred_shifts)
    can_n = sum(1 for emp in employees if "N" in emp.preferred_shifts)
    if rule.min_d > can_d:
        raise ValueError(
            f"白班 D 每日至少要 {rule.min_d} 人，但只有 {can_d} 位員工的偏好含 D（能排 D）。"
            "請降低 D 下限或讓更多人能排 D。"
        )
    if rule.min_e > can_e:
        raise ValueError(
            f"小夜 E 每日至少要 {rule.min_e} 人，但只有 {can_e} 位員工的偏好含 E（能排 E）。"
            "請降低 E 下限或讓更多人能排 E。"
        )
    if rule.min_n > can_n:
        raise ValueError(
            f"大夜 N 每日至少要 {rule.min_n} 人，但只有 {can_n} 位員工的偏好含 N（能排 N）。"
            "請降低 N 下限或讓更多人能排 N。"
        )


def parse_previous_last_shift(raw: str) -> tuple[int, str]:
    text = (raw or "").strip().upper()
    if not text:
        return 1, "off"
    idx = 0
    while idx < len(text) and text[idx].isdigit():
        idx += 1
    if idx == 0:
        raise ValueError(f"上段最後一班格式錯誤：{raw}")
    count = int(text[:idx])
    token = text[idx:].strip()
    token_map = {
        "D": "D",
        "E": "E",
        "N": "N",
        "OFF": "off",
        "OF": "off",
        "R": "R",
    }
    if token not in token_map:
        raise ValueError(f"上段最後一班格式錯誤：{raw}")
    return count, token_map[token]


def solve_schedule(employees: Sequence[Employee], days: Sequence[date], rule: RuleConfig) -> Dict[str, Dict[date, str]]:
    if len(days) == 0:
        raise ValueError("排班區間不可為空。")
    _preflight_staffing(employees, rule)

    holiday_dates = {d for d in days if d.weekday() >= 5}
    if rule.national_holidays:
        holiday_dates.update(d for d in rule.national_holidays if d in set(days))

    model = cp_model.CpModel()
    e_size = len(employees)
    d_size = len(days)

    x = {}
    for e in range(e_size):
        for d in range(d_size):
            for s in ALL_TYPES:
                x[(e, d, s)] = model.NewBoolVar(f"x_e{e}_d{d}_{s}")

    for e in range(e_size):
        for d in range(d_size):
            model.Add(sum(x[(e, d, s)] for s in ALL_TYPES) == 1)

    day_index = {day: i for i, day in enumerate(days)}
    for e, emp in enumerate(employees):
        requested_set = set(emp.requested_r_dates)
        for rd in emp.requested_r_dates:
            if rd not in day_index:
                raise ValueError(f"{emp.name} 的 R 日期 {rd} 不在排班區間")
            model.Add(x[(e, day_index[rd], "R")] == 1)
        # 非指定 R 日期不可排 R，休假時一律使用 off
        for d, day in enumerate(days):
            if day not in requested_set:
                model.Add(x[(e, d, "R")] == 0)

    for e, emp in enumerate(employees):
        for d in range(d_size):
            for s in SHIFT_TYPES:
                if s not in emp.preferred_shifts:
                    model.Add(x[(e, d, s)] == 0)

    # 跨曆日銜接：兩日合併視為 D→E→N→D→E→N 的片段，禁止「中間只隔一班」的逆向跳法。
    # 禁止：E 次日 D；N 次日 D；N 次日 E。（D 次日 E、E 次日 N、D 次日 N 等仍允許。）
    for e in range(e_size):
        for d in range(d_size - 1):
            model.Add(x[(e, d, "E")] + x[(e, d + 1, "D")] <= 1)
            model.Add(x[(e, d, "N")] + x[(e, d + 1, "D")] <= 1)
            model.Add(x[(e, d, "N")] + x[(e, d + 1, "E")] <= 1)

    for d in range(d_size):
        model.Add(sum(x[(e, d, "D")] for e in range(e_size)) >= rule.min_d)
        model.Add(sum(x[(e, d, "D")] for e in range(e_size)) <= rule.max_d)
        model.Add(sum(x[(e, d, "E")] for e in range(e_size)) >= rule.min_e)
        model.Add(sum(x[(e, d, "E")] for e in range(e_size)) <= rule.max_e)
        model.Add(sum(x[(e, d, "N")] for e in range(e_size)) >= rule.min_n)
        model.Add(sum(x[(e, d, "N")] for e in range(e_size)) <= rule.max_n)

    penalty_overrun = []
    weekend_rest_vars = []
    for e, emp in enumerate(employees):
        model.Add(sum(x[(e, d, "R")] + x[(e, d, "off")] for d in range(d_size)) >= emp.min_rest_days)
        hard_limit = emp.max_consecutive_work_days + (1 if emp.allow_fifth_day else 0)
        work = [sum(x[(e, d, s)] for s in SHIFT_TYPES) for d in range(d_size)]

        prev_count, prev_type = parse_previous_last_shift(emp.previous_last_shift)
        prev_work_streak = prev_count if prev_type in SHIFT_TYPES else 0

        # 是否上班（D/E/N 任一）
        w = [model.NewBoolVar(f"w_e{e}_d{d}") for d in range(d_size)]
        for d in range(d_size):
            model.Add(work[d] == w[d])

        # 連續上班天數（含上段銜接；遇到 off/R 歸零）
        streak = [model.NewIntVar(0, hard_limit, f"streak_e{e}_d{d}") for d in range(d_size)]
        if d_size > 0:
            model.Add(streak[0] == 0).OnlyEnforceIf(w[0].Not())
            model.Add(streak[0] == prev_work_streak + 1).OnlyEnforceIf(w[0])
            for d in range(1, d_size):
                model.Add(streak[d] == 0).OnlyEnforceIf(w[d].Not())
                both_work = model.NewBoolVar(f"both_work_e{e}_d{d}")
                model.AddMultiplicationEquality(both_work, [w[d], w[d - 1]])
                model.Add(streak[d] == streak[d - 1] + 1).OnlyEnforceIf(both_work)
                only_today = model.NewBoolVar(f"only_today_e{e}_d{d}")
                model.Add(only_today <= w[d])
                model.Add(only_today + w[d - 1] <= 1)
                model.Add(only_today >= w[d] - w[d - 1])
                model.Add(streak[d] == 1).OnlyEnforceIf(only_today)

        for d in range(d_size):
            model.Add(streak[d] <= hard_limit)

        # 上段最後為 E：本段第一個上班日不可 D；為 N：第一個上班日不可 D 或 E（與上列跨日規則一致）
        if prev_type == "E":
            for d in range(d_size):
                sum_before = sum(w[i] for i in range(d))
                first_work = model.NewBoolVar(f"first_work_e{e}_d{d}")
                model.Add(first_work <= w[d])
                model.Add(first_work + sum_before <= 1)
                model.Add(first_work >= w[d] - sum_before)
                model.Add(x[(e, d, "D")] == 0).OnlyEnforceIf(first_work)
        elif prev_type == "N":
            for d in range(d_size):
                sum_before = sum(w[i] for i in range(d))
                first_work = model.NewBoolVar(f"first_work_e{e}_d{d}")
                model.Add(first_work <= w[d])
                model.Add(first_work + sum_before <= 1)
                model.Add(first_work >= w[d] - sum_before)
                model.Add(x[(e, d, "D")] == 0).OnlyEnforceIf(first_work)
                model.Add(x[(e, d, "E")] == 0).OnlyEnforceIf(first_work)

        weekend_rest = model.NewIntVar(0, d_size, f"weekend_rest_e{e}")
        model.Add(
            weekend_rest
            == sum(
                x[(e, d, "R")] + x[(e, d, "off")]
                for d, day in enumerate(days)
                if day in holiday_dates
            )
        )
        weekend_rest_vars.append(weekend_rest)

        for d in range(d_size):
            hit_limit = model.NewBoolVar(f"hit_limit_e{e}_d{d}")
            model.Add(streak[d] == hard_limit).OnlyEnforceIf(hit_limit)
            model.Add(streak[d] <= hard_limit - 1).OnlyEnforceIf(hit_limit.Not())
            penalty_overrun.append(hit_limit)

    # 上段夜班數 + 本段該曆月至月底前 E/N ≥ 門檻（可超過）；月底以日曆月最後一天為準
    first_year = days[0].year
    first_month = days[0].month
    month_end_day = calendar.monthrange(first_year, first_month)[1]
    month_end_date = date(first_year, first_month, month_end_day)
    if month_end_date in set(days):
        month_end_idx = days.index(month_end_date)
        for e, emp in enumerate(employees):
            if emp.previous_night_count is None:
                continue

            # 先做可行性檢查，避免無解時只看到模糊錯誤
            month_days = [days[d] for d in range(month_end_idx + 1) if days[d].year == first_year and days[d].month == first_month]
            requested_r_set = {d for d in emp.requested_r_dates if d in set(month_days)}
            can_take_night_shift = bool(set(emp.preferred_shifts) & {"E", "N"})
            if can_take_night_shift:
                max_current_nights = len(month_days) - len(requested_r_set)
            else:
                max_current_nights = 0
            max_total = emp.previous_night_count + max_current_nights
            if max_total < NIGHT_COUNT_MONTH_MIN:
                raise ValueError(
                    f"{emp.name} 在月底前最多只能湊到 {max_total} 班夜班(上段{emp.previous_night_count}+本段最多{max_current_nights})，"
                    f"無法達到至少 {NIGHT_COUNT_MONTH_MIN} 班，請調整上段夜班數、偏好班別或排班區間。"
                )

            current_nights = sum(
                x[(e, d, "E")] + x[(e, d, "N")]
                for d in range(month_end_idx + 1)
                if days[d].year == first_year and days[d].month == first_month
            )
            model.Add(current_nights + emp.previous_night_count >= NIGHT_COUNT_MONTH_MIN)

    off_count = sum(x[(e, d, "off")] for e in range(e_size) for d in range(d_size))
    preference_score = sum(
        emp.preferred_shift_weights.get(s, 0) * x[(e, d, s)]
        for e, emp in enumerate(employees)
        for d in range(d_size)
        for s in emp.preferred_shifts
    )
    weekend_rest_max = model.NewIntVar(0, d_size, "weekend_rest_max")
    weekend_rest_min = model.NewIntVar(0, d_size, "weekend_rest_min")
    for v in weekend_rest_vars:
        model.Add(v <= weekend_rest_max)
        model.Add(v >= weekend_rest_min)
    weekend_rest_gap = model.NewIntVar(0, d_size, "weekend_rest_gap")
    model.Add(weekend_rest_gap == weekend_rest_max - weekend_rest_min)

    # 先滿足可行性，再偏好：假日休假平均、少 off、少連5天、高偏好分數
    model.Minimize(weekend_rest_gap * 1000 + off_count * 100 + sum(penalty_overrun) * 20 - preference_score)

    solver = cp_model.CpSolver()
    solver.parameters.random_seed = rule.random_seed
    # 複雜約束下 20s 常得到 UNKNOWN（非證明無解）；略放寬時間較易找到可行解
    solver.parameters.max_time_in_seconds = 180
    status = solver.Solve(model)
    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        label = solver.StatusName(status)
        if status == cp_model.INFEASIBLE:
            raise RuntimeError(
                f"求解器判定條件無解（{label}）："
                "每日人力下限、預假、連續上班、跨日班序（禁 E→D、N→D、N→E）、或「月底前夜班數至少 20（可超過）」等無法同時滿足。"
                "請檢查：預假是否過多或過集中、上段最後一班與夜班累計、上段末班若為 E／N 是否與本段衝突。"
            )
        if status == cp_model.UNKNOWN:
            raise RuntimeError(
                f"在計算時間內尚未找到可行解（{label}），不代表必然無解。"
                "請按「下一版」換亂數種子再試，或略放寬每日 D/E/N 下限、減少預假；若仍失敗再延長區間分段排。"
            )
        raise RuntimeError(f"排班求解未完成（{label}），請放寬條件或稍後再試。")

    result: Dict[str, Dict[date, str]] = {}
    for e, emp in enumerate(employees):
        result[emp.name] = {}
        for d, day in enumerate(days):
            assigned = None
            for s in ALL_TYPES:
                if solver.Value(x[(e, d, s)]) == 1:
                    assigned = s
                    break
            if assigned is None:
                raise RuntimeError("求解結果異常")
            result[emp.name][day] = assigned
    return result


def export_to_excel(
    path: str,
    employees: Sequence[Employee],
    schedule: Dict[str, Dict[date, str]],
    days: Sequence[date],
    national_holidays: Set[date] | None = None,
) -> None:
    holiday_dates = {d for d in days if d.weekday() >= 5}
    if national_holidays:
        holiday_dates.update(d for d in national_holidays if d in set(days))

    wb = Workbook()
    ws = wb.active
    ws.title = "排班表"

    ws.cell(row=1, column=1, value="姓名")
    ws.cell(row=1, column=2, value="上段最後班")
    ws.cell(row=1, column=3, value="上段夜班")
    first_day_col = 4
    for i, day in enumerate(days, start=first_day_col):
        ws.cell(row=1, column=i, value=day.day)
        ws.cell(row=2, column=i, value=("一二三四五六日")[day.weekday()])
    ws.cell(row=1, column=first_day_col + len(days), value="休假天數")
    ws.cell(row=1, column=first_day_col + len(days) + 1, value="假日天數")

    fills = {
        "R": PatternFill(start_color="E6CCFF", end_color="E6CCFF", fill_type="solid"),
        "off": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),
        "D": PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
        "E": PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
        "N": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
    }

    for r, emp in enumerate(employees, start=3):
        ws.cell(row=r, column=1, value=emp.name)
        ws.cell(row=r, column=2, value=emp.previous_last_shift)
        ws.cell(row=r, column=3, value=format_night_carryover_display(emp))
        rest_total = 0
        holiday_rest = 0
        for c, day in enumerate(days, start=first_day_col):
            v = schedule[emp.name][day]
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = fills[v]
            if v in REST_TYPES:
                rest_total += 1
                if day in holiday_dates:
                    holiday_rest += 1
        ws.cell(row=r, column=first_day_col + len(days), value=rest_total)
        ws.cell(row=r, column=first_day_col + len(days) + 1, value=holiday_rest)

    summary_row = len(employees) + 5
    ws.cell(row=summary_row, column=1, value="班別統計")
    ws.cell(row=summary_row + 1, column=1, value="白班 D")
    ws.cell(row=summary_row + 2, column=1, value="小夜 E")
    ws.cell(row=summary_row + 3, column=1, value="大夜 N")
    for c, day in enumerate(days, start=first_day_col):
        ws.cell(row=summary_row, column=c, value=day.day)
        ws.cell(row=summary_row + 1, column=c, value=sum(1 for emp in employees if schedule[emp.name][day] == "D"))
        ws.cell(row=summary_row + 2, column=c, value=sum(1 for emp in employees if schedule[emp.name][day] == "E"))
        ws.cell(row=summary_row + 3, column=c, value=sum(1 for emp in employees if schedule[emp.name][day] == "N"))

    total_cols = first_day_col + len(days) + 1
    for c in range(1, total_cols + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for c in range(first_day_col, first_day_col + len(days)):
        ws.column_dimensions[get_column_letter(c)].width = 5
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions[get_column_letter(first_day_col + len(days))].width = 10
    ws.column_dimensions[get_column_letter(first_day_col + len(days) + 1)].width = 10

    for row in ws.iter_rows(min_row=1, max_row=summary_row + 3, min_col=1, max_col=total_cols):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    wb.save(path)


def load_config(path: str):
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    employees = build_employees(data["employees"])
    days = date_range(parse_date(data["date_range"]["start"]), parse_date(data["date_range"]["end"]))
    e_cfg = data["daily_requirements"]["E"]
    n_cfg = data["daily_requirements"]["N"]
    e_min = e_cfg["min"] if "min" in e_cfg else e_cfg["required"]
    e_max = e_cfg["max"] if "max" in e_cfg else e_cfg["required"]
    n_min = n_cfg["min"] if "min" in n_cfg else n_cfg["required"]
    n_max = n_cfg["max"] if "max" in n_cfg else n_cfg["required"]

    rule = RuleConfig(
        min_d=data["daily_requirements"]["D"]["min"],
        max_d=data["daily_requirements"]["D"]["max"],
        min_e=e_min,
        max_e=e_max,
        min_n=n_min,
        max_n=n_max,
        national_holidays={parse_date(d) for d in data.get("national_holidays", [])},
        random_seed=data.get("random_seed", 42),
    )
    return employees, days, rule, data.get("output_file", "nurse_schedule.xlsx")


def main() -> None:
    parser = argparse.ArgumentParser(description="護理排班程式")
    parser.add_argument("--config", default="config.sample.json", help="JSON 設定檔路徑")
    args = parser.parse_args()

    employees, days, rule, output = load_config(args.config)
    schedule = solve_schedule(employees, days, rule)
    export_to_excel(output, employees, schedule, days, rule.national_holidays)
    print(f"排班完成，已輸出：{output}")


if __name__ == "__main__":
    main()
